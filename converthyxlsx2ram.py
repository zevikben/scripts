#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
convert_income.py
==================
Converts an input "income" xlsx file into an accounting-journal XLS file
(legacy binary .xls format), using "income_base_template.xls" as the
target template.

Base directory: c:\\veka\\scr
    - the template file "income_base_template.xls" must live there
    - relative input/output paths given on the command line are resolved
      against this directory too

NOTE ON FILE FORMATS
    - The TEMPLATE and OUTPUT are now genuine legacy .xls files (BIFF
      format), read/written with xlrd + xlwt + xlutils.
    - The INPUT income file is still read as modern .xlsx via openpyxl -
      only the template/output side was asked to change to .xls.

Template rules
    - Rows 1-4 of the template are NEVER touched (headers / other content).
      xlutils.copy() carries every existing row/style from the template
      into the new workbook untouched; we only write to rows 5+.
    - New data rows are written starting at row 5.
    - Output columns are located by header text, searched in rows 1-4 of
      the template, so the exact column position in the template doesn't
      matter as long as the header text matches.

Column mapping (input -> output)
    מספר המסמך      -> אסמכתא
    לקוח            -> תאור
    תאריך יצירה      -> תאריך            (parsed and written as a real date, formatted as date)
    תאריך ערך        -> תאריך ערך         (parsed and written as a real date, formatted as date)
    (constant)       -> קוד חשבון חובה  = 21001, every row
    (constant)       -> קוד חשבון זכות  = 71001, every row
    (parameter)      -> מספר מנה        = the batch number given on the command line

Per input row, up to two output rows are produced:
    a. if סכום חייב מע"מ != 0:
           סך פקודה = סכום חייב מע"מ + מע"מ,  סוג פקודה = "חמ77"
    b/c. if סכום פטור מע"מ != 0:
           סך פקודה = סכום פטור מע"מ,          סוג פקודה = "חמפחדש"
       (this is the SAME rule whether or not rule (a) also fired for this
        input row - if it did, this becomes the "duplicate" second row
        described in requirement c; if rule (a) did not fire, this is the
        row's only output line - requirement b)
    If both amounts are 0, no output row is produced for that input row.

Usage
    python convert_income.py <input.xlsx> <output.xls> <batch_number>

Example
    python convert_income.py incomes_2026_01.xlsx journal_2026_01.xls 12

Dependencies
    pip install openpyxl xlrd xlwt xlutils
"""

import sys
import os
import datetime

import openpyxl
import xlrd
import xlwt
from xlutils.copy import copy as xl_copy

# ---------------------------------------------------------------------------
# Configuration - adjust here if the template's location/name ever changes.
# ---------------------------------------------------------------------------
BASE_DIR = r"c:\veka\scr"
TEMPLATE_NAME = "income_base_template.xls"

DATA_START_ROW = 5  # 1-indexed, matching the original spec: rows 1-4 untouched

DEBIT_ACCOUNT_CODE = 21001
CREDIT_ACCOUNT_CODE = 71001

ORDER_TYPE_VAT = "חמ77"
ORDER_TYPE_EXEMPT = "חמפחדש"

DEFAULT_DATE_FORMAT = "DD/MM/YYYY"
DATE_STRING_FORMATS = (
    "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y",
    "%Y-%m-%d", "%Y/%m/%d",
    "%d/%m/%y", "%d-%m-%y",
)


def parse_date_value(value):
    """Return value as a real datetime.date/datetime, so it's written to the
    output as an actual date (not text that merely looks like one). Accepts
    values that are already date/datetime objects, or common textual date
    formats. Returns None if empty/unparseable so the caller can decide."""
    if value is None or value == "":
        return None
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value
    text = str(value).strip()
    for fmt in DATE_STRING_FORMATS:
        try:
            return datetime.datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None  # unparseable - caller decides how to handle


INPUT_HEADERS = {
    "doc_no": 'מספר המסמך',
    "client": 'לקוח',
    "create_date": 'תאריך יצירה',
    "value_date": 'תאריך ערך',
    "vat_liable_amount": 'סכום חייב מע"מ',
    "vat_amount": 'מע"מ',
    "vat_exempt_amount": 'סכום פטור מע"מ',
}

OUTPUT_HEADERS = {
    "asmachta": 'אסמכתא',
    "description": 'תאור',
    "date": 'תאריך',
    "value_date": 'תאריך ערך',
    "debit_account": 'קוד חשבון חובה',
    "credit_account": 'קוד חשבון זכות',
    "batch_no": 'מספר מנה',
    "order_sum": 'סך פקודה',
    "order_type": 'סוג פקודה',
}


# ---------------------------------------------------------------------------
# .xlsx helpers (input file only)
# ---------------------------------------------------------------------------
def find_header_columns_xlsx(ws, header_names, search_rows):
    """1-indexed rows/cols (openpyxl convention). Returns (mapping, header_row)."""
    mapping = {}
    found_row = None
    for row in search_rows:
        if row > ws.max_row:
            continue
        for cell in ws[row]:
            if cell.value is None:
                continue
            text = str(cell.value).strip()
            for key, name in header_names.items():
                if text == name and key not in mapping:
                    mapping[key] = cell.column
                    found_row = row
    missing = [name for key, name in header_names.items() if key not in mapping]
    if missing:
        raise ValueError(
            "Could not find the following expected header(s) in the INPUT file: "
            + ", ".join(missing)
            + f" (searched rows {list(search_rows)})"
        )
    return mapping, found_row


# ---------------------------------------------------------------------------
# .xls helpers (template / output file)
# ---------------------------------------------------------------------------
def find_header_columns_xls(sheet, header_names, search_rows_0idx):
    """0-indexed rows/cols (xlrd convention). Returns (mapping, header_row_0idx)."""
    mapping = {}
    found_row = None
    for row in search_rows_0idx:
        if row >= sheet.nrows:
            continue
        for col in range(sheet.ncols):
            val = sheet.cell_value(row, col)
            if val is None or val == "":
                continue
            text = str(val).strip()
            for key, name in header_names.items():
                if text == name and key not in mapping:
                    mapping[key] = col
                    found_row = row
    missing = [name for key, name in header_names.items() if key not in mapping]
    if missing:
        raise ValueError(
            "Could not find the following expected header(s) in the TEMPLATE: "
            + ", ".join(missing)
            + f" (searched 0-indexed rows {list(search_rows_0idx)})"
        )
    return mapping, found_row


def resolve_path(path_arg):
    """Resolve a path against BASE_DIR unless it's already absolute."""
    return path_arg if os.path.isabs(path_arg) else os.path.join(BASE_DIR, path_arg)


def parse_batch_no(raw):
    """Keep the batch number numeric if it looks numeric, else keep it as text."""
    try:
        return int(raw)
    except ValueError:
        return raw


def main():
    if len(sys.argv) != 4:
        print("Usage: python convert_income.py <input.xlsx> <output.xls> <batch_number>")
        sys.exit(1)

    input_arg, output_arg, batch_arg = sys.argv[1], sys.argv[2], sys.argv[3]
    batch_no = parse_batch_no(batch_arg)

    if os.path.isdir(BASE_DIR):
        os.chdir(BASE_DIR)

    input_path = resolve_path(input_arg)
    output_path = resolve_path(output_arg)
    template_path = resolve_path(TEMPLATE_NAME)

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template file not found: {template_path}")
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Input file not found: {input_path}")

    # --- open the .xls template with xlrd (formatting_info=True is required
    #     for xlutils.copy to be able to carry the existing styles over) ---
    rb = xlrd.open_workbook(template_path, formatting_info=True)
    template_sheet = rb.sheet_by_index(0)

    out_cols, _out_header_row = find_header_columns_xls(
        template_sheet, OUTPUT_HEADERS, search_rows_0idx=range(0, 4)
    )  # rows 1-4 (1-indexed) == 0-3 (0-indexed)

    # --- xlutils.copy() gives back an xlwt.Workbook pre-populated with every
    #     existing row/style from the template - rows 1-4 need no further
    #     action at all, they're already carried over untouched ---
    out_wb = xl_copy(rb)
    out_ws = out_wb.get_sheet(0)

    date_style = xlwt.XFStyle()
    date_style.num_format_str = DEFAULT_DATE_FORMAT
    plain_style = xlwt.XFStyle()

    # --- read the input workbook (still modern .xlsx via openpyxl) ---
    in_wb = openpyxl.load_workbook(input_path, data_only=True)
    in_ws = in_wb.active
    in_cols, in_header_row = find_header_columns_xlsx(in_ws, INPUT_HEADERS, search_rows=range(1, 4))

    max_row = in_ws.max_row
    out_row_1idx = DATA_START_ROW
    rows_from_rule_a = 0
    rows_from_rule_bc = 0
    skipped = 0
    unparseable_dates = 0

    def gv(row, key):
        return in_ws.cell(row=row, column=in_cols[key]).value

    def write_common(out_row_0idx, doc_no, client, create_date, value_date):
        out_ws.write(out_row_0idx, out_cols["asmachta"], doc_no, plain_style)
        out_ws.write(out_row_0idx, out_cols["description"], client, plain_style)
        out_ws.write(out_row_0idx, out_cols["date"], create_date, date_style)
        out_ws.write(out_row_0idx, out_cols["value_date"], value_date, date_style)
        out_ws.write(out_row_0idx, out_cols["debit_account"], DEBIT_ACCOUNT_CODE, plain_style)
        out_ws.write(out_row_0idx, out_cols["credit_account"], CREDIT_ACCOUNT_CODE, plain_style)
        out_ws.write(out_row_0idx, out_cols["batch_no"], batch_no, plain_style)

    for r in range(in_header_row + 1, max_row + 1):
        doc_no = gv(r, "doc_no")
        if doc_no is None:
            continue  # blank row - skip

        client = gv(r, "client")

        create_date_raw = gv(r, "create_date")
        value_date_raw = gv(r, "value_date")
        create_date = parse_date_value(create_date_raw)
        value_date = parse_date_value(value_date_raw)
        if create_date_raw not in (None, "") and create_date is None:
            unparseable_dates += 1
            create_date = create_date_raw  # fall back to the raw value rather than losing it
        if value_date_raw not in (None, "") and value_date is None:
            unparseable_dates += 1
            value_date = value_date_raw

        vat_liable = gv(r, "vat_liable_amount") or 0
        vat_amount = gv(r, "vat_amount") or 0
        vat_exempt = gv(r, "vat_exempt_amount") or 0

        row_written = False
        out_row_0idx = out_row_1idx - 1  # xlwt is 0-indexed

        # rule (a): סכום חייב מע"מ != 0 -> חמ77 row
        if vat_liable != 0:
            write_common(out_row_0idx, doc_no, client, create_date, value_date)
            out_ws.write(out_row_0idx, out_cols["order_sum"], vat_liable + vat_amount, plain_style)
            out_ws.write(out_row_0idx, out_cols["order_type"], ORDER_TYPE_VAT, plain_style)
            out_row_1idx += 1
            rows_from_rule_a += 1
            row_written = True
            out_row_0idx = out_row_1idx - 1

        # rules (b)/(c): סכום פטור מע"מ != 0 -> חמפחדש row
        # (this is the "duplicate" second row of requirement c, when rule (a)
        # above also fired for this same input row; otherwise it's the row's
        # only output line, per requirement b)
        if vat_exempt != 0:
            write_common(out_row_0idx, doc_no, client, create_date, value_date)
            out_ws.write(out_row_0idx, out_cols["order_sum"], vat_exempt, plain_style)
            out_ws.write(out_row_0idx, out_cols["order_type"], ORDER_TYPE_EXEMPT, plain_style)
            out_row_1idx += 1
            rows_from_rule_bc += 1
            row_written = True

        if not row_written:
            skipped += 1

    out_wb.save(output_path)

    total_written = out_row_1idx - DATA_START_ROW
    print(f"Done: {output_path}")
    print(f"  rows written: {total_written}")
    print(f"    from rule a  (חמ77):     {rows_from_rule_a}")
    print(f"    from rule b/c (חמפחדש): {rows_from_rule_bc}")
    if skipped:
        print(f"  skipped {skipped} input row(s) where both amounts were 0 (no rule matched)")
    if unparseable_dates:
        print(f"  WARNING: {unparseable_dates} date value(s) could not be parsed and were "
              f"copied through as-is (may not display as a proper date in the output)")


if __name__ == "__main__":
    main()
