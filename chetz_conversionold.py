# -*- coding: utf-8 -*-
"""
chetz conversion
----------------
Converts rows of a given month from the bank export (xlsx) into an income
journal (xls) built from income_base_template.xls.

Usage (from c:\\veka\\scr):
    python chetz_conversion.py <month> <output_file.xls> <batch_number>

    month         MM/YYYY (e.g. 09/2026) or just MM (any year)
    output_file   file name, written to c:\\veka\\datafiles
    batch_number  value written to column "מספר מנה" on every row

Example:
    python chetz_conversion.py 09/2026 income_09_2026.xls 125

Requires:  pip install openpyxl xlrd xlwt xlutils
"""
import os
import re
import sys
import shutil
import datetime as dt

import openpyxl
from openpyxl.styles import PatternFill
import xlrd
import xlwt
from xlutils.copy import copy as xl_copy

# ---------------------------------------------------------------- paths
BASE_DIR = r"c:\veka\scr"
DATA_DIR = r"c:\veka\datafiles"
TEMPLATE = os.path.join(BASE_DIR, "income_base_template.xls")
INPUT_XLSX = os.path.join(DATA_DIR, "chetz_base_iteration.xlsx")
SUPPLIERS = os.path.join(DATA_DIR, "chetz_suppliers_list.xls")

# ---------------------------------------------------------------- constants
TRANSFER_MARK = "העברה אל:"
DESCRIPTION_PREFIX = "העברה-"
CREDIT_ACCOUNT = 11003
TEMPLATE_PROTECTED_ROWS = 4          # rows 1-4 of the template are untouched

# input columns
IN_DATE = "תאריך"
IN_VALUE_DATE = "תאריך ערך"
IN_DESC = "תיאור מורחב"
IN_REF = "אסמכתא"
IN_DEBIT = "בחובה"
IN_NOTES = "הערות"            # optional column

SALARY_WORD = "משכורת"
SALARY_ACCOUNT = 54001

# output columns
OUT_REF = "אסמכתא"
OUT_DESC = "תאור"
OUT_DATE = "תאריך"
OUT_VALUE_DATE = "תאריך ערך"
OUT_DEBIT_ACC = "קוד חשבון חובה"
OUT_CREDIT_ACC = "קוד חשבון זכות"
OUT_BATCH = "מספר מנה"
OUT_AMOUNT = "סך פקודה"

# supplier list columns
SUP_NAME = "שם חשבון"
SUP_KEY = "מפתח חשבון"

PURPLE = PatternFill("solid", start_color="D9B3FF", end_color="D9B3FF")
RED = PatternFill("solid", start_color="FF8080", end_color="FF8080")
BLUE = PatternFill("solid", start_color="9BC2E6", end_color="9BC2E6")

STATUS_COL = "סטטוס המרה"      # added to the input file, explains each color
LEGEND = [
    "Color legend:",
    "  PURPLE - converted OK (supplier found, or salary -> 54001)",
    "  BLUE   - only part of the supplier name was found - check the code",
    "  RED    - supplier not found or data missing - code left empty",
]


# ================================================================ output / log
LOG_FILE = os.path.join(DATA_DIR, "chetz_conversion_log.txt")
_log_lines = []

try:  # make the console print Hebrew as UTF-8
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass


def out(msg=""):
    _log_lines.append(msg)
    try:
        print(msg)
    except Exception:
        pass


def write_log():
    try:
        with open(LOG_FILE, "w", encoding="utf-8-sig") as f:
            f.write(dt.datetime.now().strftime("%d/%m/%Y %H:%M") + "  "
                    + " ".join(sys.argv[1:]) + "\n")
            f.write("\n".join(_log_lines) + "\n")
        print(f"Log    : {LOG_FILE}")
    except Exception:
        pass


def fail(msg):
    out("ERROR: " + msg)
    write_log()
    sys.exit(1)


# ================================================================ helpers
def norm(s):
    """Normalise a header / name for comparison."""
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()


def parse_date(value, datemode=0):
    """Return a datetime.date from a cell value (datetime, Excel serial or text)."""
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        try:
            return xlrd.xldate_as_datetime(value, datemode).date()
        except Exception:
            return None
    m = re.search(r"(\d{1,2})\D(\d{1,2})\D(\d{2,4})", str(value))
    if not m:
        return None
    d, mo, y = (int(x) for x in m.groups())
    if y < 100:
        y += 2000
    try:
        return dt.date(y, mo, d)
    except ValueError:
        return None


def parse_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return value
    s = str(value).replace(",", "").replace("₪", "").strip()
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return None


def parse_month(arg):
    """'09/2026' -> (9, 2026);  '9' -> (9, None)"""
    m = re.fullmatch(r"\s*(\d{1,2})\s*(?:[/\-.]\s*(\d{4}))?\s*", arg)
    if not m:
        fail(f"Bad month '{arg}'. Use MM/YYYY (e.g. 09/2026) or MM.")
    month = int(m.group(1))
    if not 1 <= month <= 12:
        fail(f"Bad month '{arg}'.")
    return month, (int(m.group(2)) if m.group(2) else None)


def supplier_from_description(desc):
    """Return the two words that follow 'העברה אל:' (or None)."""
    text = norm(desc)
    idx = text.find(TRANSFER_MARK)
    if idx < 0:
        return None
    words = text[idx + len(TRANSFER_MARK):].split()
    if not words:
        return ""
    return " ".join(words[:2])


def find_header_row(rows, required, max_scan=15):
    """rows: list of lists. Return (row_index, {header: col_index})."""
    for r, row in enumerate(rows[:max_scan]):
        cols = {norm(v): c for c, v in enumerate(row) if norm(v)}
        if all(h in cols for h in required):
            return r, cols
    return None, None


# ================================================================ loaders
def load_suppliers():
    book = xlrd.open_workbook(SUPPLIERS)
    sh = book.sheet_by_index(0)
    rows = [sh.row_values(r) for r in range(sh.nrows)]
    hr, cols = find_header_row(rows, [SUP_NAME, SUP_KEY])
    if hr is None:
        fail(f"Columns '{SUP_NAME}' / '{SUP_KEY}' not found in {SUPPLIERS}")
    suppliers = {}
    for row in rows[hr + 1:]:
        name = norm(row[cols[SUP_NAME]])
        key = row[cols[SUP_KEY]]
        if name and key not in ("", None):
            if isinstance(key, float) and key.is_integer():
                key = int(key)
            suppliers.setdefault(name, key)
    return suppliers


def lookup_supplier(suppliers, name):
    """
    Search supplier_name inside the full text of every "שם חשבון".
    Returns (key, status, account_name):
      status "full"    - the whole name appears in the account name
      status "partial" - only some of its words appear  (row colored blue)
      status None      - nothing found                  (row colored red)
    """
    if not name:
        return None, None, None
    # 1. whole name found inside the account text (prefer the shortest = closest)
    full = [k for k in suppliers if name in k]
    if full:
        best = min(full, key=len)
        return suppliers[best], "full", best
    # 2. part of the name: the account that contains the most of its words
    words = [w for w in name.split() if len(w) >= 2]
    best, best_score = None, 0
    for k in suppliers:
        k_words = k.split()
        score = sum(1 for w in words if w in k_words or w in k)
        if score > best_score:
            best, best_score = k, score
    if best:
        return suppliers[best], "partial", best
    return None, None, None


# ================================================================ main
def main():
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)

    month, year = parse_month(sys.argv[1])
    out_name = sys.argv[2]
    if not out_name.lower().endswith(".xls"):
        out_name += ".xls"
    out_path = os.path.join(DATA_DIR, os.path.basename(out_name))
    batch = parse_number(sys.argv[3])
    if batch is None:
        batch = sys.argv[3]

    os.chdir(BASE_DIR)
    for p in (TEMPLATE, INPUT_XLSX, SUPPLIERS):
        if not os.path.exists(p):
            fail(f"File not found: {p}")

    suppliers = load_suppliers()
    # write 54001 the same way the supplier codes are stored (text vs number),
    # so it lines up like the other codes in "קוד חשבון חובה"
    n_text = sum(isinstance(v, str) for v in suppliers.values())
    salary_acc = (str(SALARY_ACCOUNT) if n_text >= len(suppliers) / 2
                  else SALARY_ACCOUNT)

    # ---- template: locate output headers
    tbook = xlrd.open_workbook(TEMPLATE, formatting_info=True)
    tsheet = tbook.sheet_by_index(0)
    trows = [tsheet.row_values(r) for r in range(tsheet.nrows)]
    out_required = [OUT_REF, OUT_DESC, OUT_DATE, OUT_VALUE_DATE, OUT_DEBIT_ACC,
                    OUT_CREDIT_ACC, OUT_BATCH, OUT_AMOUNT]
    ohr, ocols = find_header_row(trows, out_required)
    if ohr is None:
        fail("Not all output column titles were found in the template:\n  "
                 + ", ".join(out_required))
    first_out_row = max(ohr + 1, TEMPLATE_PROTECTED_ROWS)   # 0-based

    wbook = xl_copy(tbook)
    wsheet = wbook.get_sheet(0)

    def style(date=False, color=None):
        s = xlwt.XFStyle()
        if date:
            s.num_format_str = "DD/MM/YYYY"
        if color:
            p = xlwt.Pattern()
            p.pattern = xlwt.Pattern.SOLID_PATTERN
            p.pattern_fore_colour = xlwt.Style.colour_map[color]
            s.pattern = p
        return s

    st = {(d, c): style(d, c) for d in (False, True)
          for c in (None, "red", "pale_blue")}

    # ---- input
    shutil.copy2(INPUT_XLSX, INPUT_XLSX + ".bak")
    ibook = openpyxl.load_workbook(INPUT_XLSX)
    isheet = ibook.worksheets[0]
    irows = [[c.value for c in row] for row in isheet.iter_rows()]
    in_required = [IN_DATE, IN_VALUE_DATE, IN_DESC, IN_REF, IN_DEBIT]
    ihr, icols = find_header_row(irows, in_required)
    if ihr is None:
        fail("Not all input column titles were found in the input file:\n  "
                 + ", ".join(in_required))
    # status column: reuse it on re-runs, otherwise add it after the last column
    if STATUS_COL in icols:
        status_col = icols[STATUS_COL] + 1
        max_col = isheet.max_column
    else:
        status_col = isheet.max_column + 1
        isheet.cell(row=ihr + 1, column=status_col, value=STATUS_COL)
        max_col = status_col
    if IN_NOTES not in icols:
        out(f"Note: column '{IN_NOTES}' not found in input - salary rule skipped")

    out_r = first_out_row
    ok_count = err_count = blue_count = 0
    partials = []
    problems = []

    for r in range(ihr + 1, len(irows)):
        row = irows[r]
        d = parse_date(row[icols[IN_DATE]])
        if d is None or d.month != month or (year and d.year != year):
            continue
        supplier_name = supplier_from_description(row[icols[IN_DESC]])
        if supplier_name is None:
            continue

        excel_row = r + 1
        errors = []

        ref = parse_number(row[icols[IN_REF]])
        if ref is None:
            ref = row[icols[IN_REF]]
            errors.append("חסרה אסמכתא") if ref in (None, "") else None
        vdate = parse_date(row[icols[IN_VALUE_DATE]])
        if vdate is None:
            errors.append("תאריך ערך לא תקין")
        amount = parse_number(row[icols[IN_DEBIT]])
        if amount is None:
            errors.append("חסר סכום בחובה")
        acc, match, matched_name = lookup_supplier(suppliers, supplier_name)
        notes = norm(row[icols[IN_NOTES]]) if IN_NOTES in icols else ""
        if match != "full" and SALARY_WORD in notes:
            acc, match, matched_name = salary_acc, "salary", SALARY_WORD
        if acc is None:
            errors.append(f"ספק לא נמצא: '{supplier_name}'")

        if errors:
            color = "red"
        elif match == "partial":
            color = "pale_blue"
        else:
            color = None
        values = {
            OUT_REF: (ref, False),
            OUT_DESC: (DESCRIPTION_PREFIX + supplier_name, False),
            OUT_DATE: (d, True),
            OUT_VALUE_DATE: (vdate, True),
            OUT_DEBIT_ACC: (acc, False),
            OUT_CREDIT_ACC: (CREDIT_ACCOUNT, False),
            OUT_BATCH: (batch, False),
            OUT_AMOUNT: (amount, False),
        }
        if color:   # paint the whole used width of the output row
            for c in range(len(trows[ohr])):
                wsheet.write(out_r, c, "", st[(False, color)])
        for title, (val, is_date) in values.items():
            wsheet.write(out_r, ocols[title], "" if val is None else val,
                         st[(is_date and val is not None, color)])
        out_r += 1

        fill = {"red": RED, "pale_blue": BLUE}.get(color, PURPLE)
        for c in range(1, max_col + 1):
            isheet.cell(row=excel_row, column=c).fill = fill
        if color == "red":
            status = "אדום - " + "; ".join(errors)
        elif color == "pale_blue":
            status = f"כחול - התאמה חלקית: '{supplier_name}' -> '{matched_name}' ({acc})"
        elif match == "salary":
            status = f"סגול - תקין (משכורת -> {acc})"
        else:
            status = f"סגול - תקין ({matched_name} -> {acc})"
        isheet.cell(row=excel_row, column=status_col, value=status)
        if color == "red":
            err_count += 1
            problems.append(f"  input row {excel_row}: " + "; ".join(errors))
        elif color == "pale_blue":
            blue_count += 1
            partials.append(f"  input row {excel_row}: '{supplier_name}' -> "
                            f"'{matched_name}' ({acc})")
        else:
            ok_count += 1

    wbook.save(out_path)
    try:
        ibook.save(INPUT_XLSX)
    except PermissionError:
        fail(f"Output written, but could not color {INPUT_XLSX} "
             "- close it in Excel and run again.")

    out(f"Output : {out_path}")
    out("\n".join(LEGEND))
    out(f"Rows   : {ok_count + blue_count + err_count} written "
        f"({ok_count} OK, {blue_count} blue - partial name, {err_count} red)")
    if partials:
        out("Partial supplier matches (blue) - please check:")
        out("\n".join(partials))
    if problems:
        out("Problems:")
        out("\n".join(problems))
    write_log()


if __name__ == "__main__":
    main()
