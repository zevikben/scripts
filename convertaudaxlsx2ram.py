#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
convert_income_3.py
====================
Converts an input "income" xlsx file (audaincome-style export) into an
accounting-journal XLS file (legacy binary .xls format), using
"income_base_template.xls" as the target template.

This is a SEPARATE program from convert_income.py and convert_income_2.py -
different input column names, a document-type filter condition, both
output date columns sourced from the same single input column, and its
own account codes / order type.

Base directory: c:\\veka\\scr
    - the template file "income_base_template.xls" must live there

Data directory: c:\\veka\\datafiles
    - relative INPUT and OUTPUT paths given on the command line are both
      resolved against this directory (created automatically for the
      output side if it doesn't exist)
    - pass an absolute path for either one to override this and use a
      different location

NOTE ON FILE FORMATS
    - The TEMPLATE and OUTPUT are genuine legacy .xls files (BIFF format),
      read/written with xlrd + xlwt + xlutils.
    - The INPUT income file is read as modern .xlsx via openpyxl.

Template rules
    - Rows 1-4 of the template are NEVER touched (headers / other content).
      xlutils.copy() carries every existing row/style from the template
      into the new workbook untouched; we only write to rows 5+.
    - New data rows are written starting at row 5.
    - Output columns are located by header text, searched in rows 1-4 of
      the template, so the exact column position in the template doesn't
      matter as long as the header text matches.

Filter condition
    Input rows are processed when סוג מסמך is one of:
      - "חשבונית מס-קבלה" (exact match) - always processed, uses the
        default קוד חשבון חובה = 21001.
      - "חשבונית מס" or "חשבונית זיכוי" - processed ONLY when שם (the
        client name) is one of two specific colleges, which also
        determines the קוד חשבון חובה used for that row instead of the
        default:
            "המכללה למנהל/רשת בתי-ספר להנדסאים"  -> קוד חשבון חובה = 21011
            "המכללה האקדמית גליל מערבי"           -> קוד חשבון חובה = 21014
        Rows with these two doc types but a different/blank שם are
        filtered out, same as any other unrecognized row.
    All other rows (e.g. subtotal/summary rows like "סך (חשבונית מס-קבלה)"
    or "סך הכנסות") are skipped entirely.

Column mapping (input -> output), for rows that pass the filter
    מספר מסמך          -> אסמכתא
    שם                 -> תאור
    תאריך הפקה          -> תאריך            (parsed and written as a real date, formatted as date)
    תאריך הפקה          -> תאריך ערך         (same source column as above, per spec)
    (constant/lookup)  -> קוד חשבון חובה  = 21001 by default, or 21011 / 21014
                          for the two special colleges above (see Filter
                          condition)
    (constant)         -> קוד חשבון זכות  = 71001, every row
    (parameter)        -> מספר מנה        = the batch number given on the command line
    סה"כ לתשלום         -> סך פקודה
    (constant)         -> סוג פקודה        = "חמ8", every row

Every input row that passes the filter produces exactly one output row -
no branching, no duplication.

Usage
    python convert_income_3.py <input.xlsx> <output.xls> <batch_number>

Example
    python convert_income_3.py audaincome_2026_06.xlsx journal_2026_06.xls 12

Dependencies
    pip install openpyxl xlrd xlwt xlutils
"""

import sys
import os
import datetime

import openpyxl
import xlrd
import xlwt
from xlutils.copy import copy as xl_copy

# ---------------------------------------------------------------------------
# Configuration - adjust here if the template's location/name ever changes.
# ---------------------------------------------------------------------------
BASE_DIR = r"c:\veka\scr"
TEMPLATE_NAME = "income_base_template.xls"
DATA_DIR = r"c:\veka\datafiles"

DATA_START_ROW = 5  # 1-indexed, matching the spec: rows 1-4 untouched

DEBIT_ACCOUNT_CODE = 21001  # default, used for DOC_TYPE_FILTER_VALUE rows
CREDIT_ACCOUNT_CODE = 71001
ORDER_TYPE = "חמ8"

DOC_TYPE_FILTER_COLUMN = 'סוג מסמך'
DOC_TYPE_FILTER_VALUE = 'חשבונית מס-קבלה'

# Additional doc types that are only processed for two specific customers,
# each with their own קוד חשבון חובה instead of the DEBIT_ACCOUNT_CODE default.
CLIENT_NAME_COLUMN = 'שם'
DOC_TYPES_WITH_CLIENT_OVERRIDE = ('חשבונית מס', 'חשבונית זיכוי')
CLIENT_DEBIT_ACCOUNT_OVERRIDES = {
    'המכללה למנהל/רשת בתי-ספר להנדסאים': 21011,
    'המכללה האקדמית גליל מערבי': 21014,
}

DEFAULT_DATE_FORMAT = "DD/MM/YYYY"
DATE_STRING_FORMATS = (
    "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y",
    "%Y-%m-%d", "%Y/%m/%d",
    "%d/%m/%y", "%d-%m-%y",
)


def parse_date_value(value):
    """Return value as a real datetime.date/datetime, so it's written to the
    output as an actual date (not text that merely looks like one). Accepts
    values that are already date/datetime objects, or common textual date
    formats. Returns None if empty/unparseable so the caller can decide."""
    if value is None or value == "":
        return None
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value
    text = str(value).strip()
    for fmt in DATE_STRING_FORMATS:
        try:
            return datetime.datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None  # unparseable - caller decides how to handle


def fix_swapped_native_date(value):
    """Known bug in the audaincome exporter: for some rows, תאריך הפקה is
    written as a genuine Excel date object (not text) with the day and
    month swapped - this happens specifically whenever the intended day
    (1-12) is also a valid month number, so Excel silently accepts the
    swapped value as a real date instead of erroring out. Text-format
    dates in the same column are NOT affected by this bug and are left
    alone (parse_date_value already handles those correctly).

    Confirmed against a real export: the affected rows form a perfect
    day-by-day sequence within the batch's actual month once swapped back
    (e.g. stored Jan7/Feb7/.../Dec7 -> swapped July1/July2/.../July12),
    while unaffected rows (day > 12, always stored as text) were already
    correct. Only native datetime/date values with day <= 12 are swapped;
    everything else (including all text dates) passes through unchanged.
    """
    if isinstance(value, (datetime.date, datetime.datetime)) and value.day <= 12:
        return value.replace(day=value.month, month=value.day)
    return value


INPUT_HEADERS = {
    "doc_type": DOC_TYPE_FILTER_COLUMN,
    "doc_no": 'מספר מסמך',
    "client": 'שם',
    "issue_date": 'תאריך הפקה',
    "total_payable": 'סה"כ לתשלום',
}

OUTPUT_HEADERS = {
    "asmachta": 'אסמכתא',
    "description": 'תאור',
    "date": 'תאריך',
    "value_date": 'תאריך ערך',
    "debit_account": 'קוד חשבון חובה',
    "credit_account": 'קוד חשבון זכות',
    "batch_no": 'מספר מנה',
    "order_sum": 'סך פקודה',
    "order_type": 'סוג פקודה',
}


# ---------------------------------------------------------------------------
# .xlsx helpers (input file only)
# ---------------------------------------------------------------------------
def find_header_columns_xlsx(ws, header_names, search_rows):
    """1-indexed rows/cols (openpyxl convention). Returns (mapping, header_row)."""
    mapping = {}
    found_row = None
    for row in search_rows:
        if row > ws.max_row:
            continue
        for cell in ws[row]:
            if cell.value is None:
                continue
            text = str(cell.value).strip()
            for key, name in header_names.items():
                if text == name and key not in mapping:
                    mapping[key] = cell.column
                    found_row = row
    missing = [name for key, name in header_names.items() if key not in mapping]
    if missing:
        raise ValueError(
            "Could not find the following expected header(s) in the INPUT file: "
            + ", ".join(missing)
            + f" (searched rows {list(search_rows)})"
        )
    return mapping, found_row


# ---------------------------------------------------------------------------
# .xls helpers (template / output file)
# ---------------------------------------------------------------------------
def find_header_columns_xls(sheet, header_names, search_rows_0idx):
    """0-indexed rows/cols (xlrd convention). Returns (mapping, header_row_0idx)."""
    mapping = {}
    found_row = None
    for row in search_rows_0idx:
        if row >= sheet.nrows:
            continue
        for col in range(sheet.ncols):
            val = sheet.cell_value(row, col)
            if val is None or val == "":
                continue
            text = str(val).strip()
            for key, name in header_names.items():
                if text == name and key not in mapping:
                    mapping[key] = col
                    found_row = row
    missing = [name for key, name in header_names.items() if key not in mapping]
    if missing:
        raise ValueError(
            "Could not find the following expected header(s) in the TEMPLATE: "
            + ", ".join(missing)
            + f" (searched 0-indexed rows {list(search_rows_0idx)})"
        )
    return mapping, found_row


def resolve_path(path_arg):
    """Resolve a path against BASE_DIR unless it's already absolute."""
    return path_arg if os.path.isabs(path_arg) else os.path.join(BASE_DIR, path_arg)


def resolve_data_path(path_arg, create_dir=False):
    """Resolve an input/output data file against DATA_DIR (c:\\veka\\datafiles)
    unless it's already absolute. If create_dir is True, DATA_DIR is created
    when missing (used for the output file - we're about to write there).
    For the input file we leave create_dir False, since a missing DATA_DIR
    in that case just means the input file itself won't be found, which is
    reported clearly by the existence check right after."""
    if os.path.isabs(path_arg):
        return path_arg
    if create_dir:
        os.makedirs(DATA_DIR, exist_ok=True)
    return os.path.join(DATA_DIR, path_arg)


def parse_batch_no(raw):
    """Keep the batch number numeric if it looks numeric, else keep it as text."""
    try:
        return int(raw)
    except ValueError:
        return raw


def main():
    if len(sys.argv) != 4:
        print("Usage: python convert_income_3.py <input.xlsx> <output.xls> <batch_number>")
        sys.exit(1)

    input_arg, output_arg, batch_arg = sys.argv[1], sys.argv[2], sys.argv[3]
    batch_no = parse_batch_no(batch_arg)

    if os.path.isdir(BASE_DIR):
        os.chdir(BASE_DIR)

    input_path = resolve_data_path(input_arg)
    output_path = resolve_data_path(output_arg, create_dir=True)
    template_path = resolve_path(TEMPLATE_NAME)

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template file not found: {template_path}")
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Input file not found: {input_path}")

    # --- open the .xls template with xlrd (formatting_info=True is required
    #     for xlutils.copy to be able to carry the existing styles over) ---
    rb = xlrd.open_workbook(template_path, formatting_info=True)
    template_sheet = rb.sheet_by_index(0)

    out_cols, _out_header_row = find_header_columns_xls(
        template_sheet, OUTPUT_HEADERS, search_rows_0idx=range(0, 4)
    )  # rows 1-4 (1-indexed) == 0-3 (0-indexed)

    # --- xlutils.copy() gives back an xlwt.Workbook pre-populated with every
    #     existing row/style from the template - rows 1-4 need no further
    #     action at all, they're already carried over untouched ---
    out_wb = xl_copy(rb)
    out_ws = out_wb.get_sheet(0)

    # Row 5 in the template holds SAMPLE/EXAMPLE data across every column
    # (not just the ones this script writes to). Only rows 1-4 are protected -
    # row 5 is a real data row and must start clean, otherwise any column
    # this script doesn't explicitly populate (e.g. מספר פקודה, אסמכתא 2)
    # would silently keep the template's stray example value.
    blank_style = xlwt.XFStyle()
    for c in range(template_sheet.ncols):
        out_ws.write(DATA_START_ROW - 1, c, "", blank_style)

    date_style = xlwt.XFStyle()
    date_style.num_format_str = DEFAULT_DATE_FORMAT
    plain_style = xlwt.XFStyle()

    # --- read the input workbook ---
    in_wb = openpyxl.load_workbook(input_path, data_only=True)
    in_ws = in_wb.active
    in_cols, in_header_row = find_header_columns_xlsx(in_ws, INPUT_HEADERS, search_rows=range(1, 4))

    max_row = in_ws.max_row
    out_row_1idx = DATA_START_ROW
    rows_written = 0
    blank_rows = []      # row numbers skipped: no סוג מסמך at all
    filtered_rows = []   # (row number, actual סוג מסמך value, reason or None) skipped
    unparseable_dates = 0

    def gv(row, key):
        return in_ws.cell(row=row, column=in_cols[key]).value

    for r in range(in_header_row + 1, max_row + 1):
        doc_type = gv(r, "doc_type")
        if doc_type is None:
            blank_rows.append(r)
            continue  # blank row - skip silently

        doc_type_text = str(doc_type).strip()
        client = gv(r, "client")
        client_text = str(client).strip() if client is not None else ""

        # --- filter: which doc types are processed, and which קוד חשבון
        #     חובה applies to this row ---
        if doc_type_text == DOC_TYPE_FILTER_VALUE:
            debit_account_code = DEBIT_ACCOUNT_CODE
        elif doc_type_text in DOC_TYPES_WITH_CLIENT_OVERRIDE:
            if client_text in CLIENT_DEBIT_ACCOUNT_OVERRIDES:
                debit_account_code = CLIENT_DEBIT_ACCOUNT_OVERRIDES[client_text]
            else:
                reason = (f"סוג מסמך '{doc_type_text}' אך {CLIENT_NAME_COLUMN} "
                           f"'{client_text or '(ריק)'}' אינו אחד משני הלקוחות המוכרים עבור סוג מסמך זה")
                filtered_rows.append((r, doc_type_text, reason))
                continue
        else:
            filtered_rows.append((r, doc_type_text, None))
            continue

        doc_no = gv(r, "doc_no")
        total_payable = gv(r, "total_payable")

        issue_date_raw = gv(r, "issue_date")
        issue_date_raw = fix_swapped_native_date(issue_date_raw)
        issue_date = parse_date_value(issue_date_raw)
        if issue_date_raw not in (None, "") and issue_date is None:
            unparseable_dates += 1
            issue_date = issue_date_raw  # fall back to the raw value rather than losing it

        out_row_0idx = out_row_1idx - 1  # xlwt is 0-indexed

        out_ws.write(out_row_0idx, out_cols["asmachta"], doc_no, plain_style)
        out_ws.write(out_row_0idx, out_cols["description"], client, plain_style)
        out_ws.write(out_row_0idx, out_cols["date"], issue_date, date_style)
        out_ws.write(out_row_0idx, out_cols["value_date"], issue_date, date_style)  # same source column
        out_ws.write(out_row_0idx, out_cols["debit_account"], debit_account_code, plain_style)
        out_ws.write(out_row_0idx, out_cols["credit_account"], CREDIT_ACCOUNT_CODE, plain_style)
        out_ws.write(out_row_0idx, out_cols["batch_no"], batch_no, plain_style)
        out_ws.write(out_row_0idx, out_cols["order_sum"], total_payable, plain_style)
        out_ws.write(out_row_0idx, out_cols["order_type"], ORDER_TYPE, plain_style)

        out_row_1idx += 1
        rows_written += 1

    out_wb.save(output_path)

    print(f"Done: {output_path}")
    print(f"  rows written: {rows_written}")

    print(f"  filtered out (doc type / client not recognized): {len(filtered_rows)}")
    if filtered_rows:
        for row, actual_value, reason in filtered_rows:
            shown_value = actual_value if actual_value else "(empty)"
            if reason:
                print(f"    row {row}: {DOC_TYPE_FILTER_COLUMN} = '{shown_value}' - {reason}")
            else:
                print(f"    row {row}: {DOC_TYPE_FILTER_COLUMN} = '{shown_value}'")

    if blank_rows:
        print(f"  skipped {len(blank_rows)} row(s) with no {DOC_TYPE_FILTER_COLUMN} value at all "
              f"(blank row):")
        print("    rows: " + ", ".join(str(r) for r in blank_rows))

    if unparseable_dates:
        print(f"  WARNING: {unparseable_dates} date value(s) could not be parsed and were "
              f"copied through as-is (may not display as a proper date in the output)")


if __name__ == "__main__":
    main()
